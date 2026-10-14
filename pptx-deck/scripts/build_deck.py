#!/usr/bin/env python3
"""Build a .pptx from a JSON deck spec. Reads the spec and any local images it
names, writes one file, uses no network. Needs: python3 -m pip install python-pptx

usage: build_deck.py spec.json out.pptx
"""
import json
import os
import sys

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt
except ImportError:
    sys.exit("python-pptx is not installed: python3 -m pip install python-pptx")

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.8)
MAX_BULLETS, MAX_WORDS, MAX_HEADLINE = 5, 14, 12
warnings = []


def rgb(hex_):
    return RGBColor.from_string(hex_.lstrip("#").upper())


def words(text):
    return len(str(text).split())


class Deck:
    def __init__(self, theme):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.font = theme.get("font", "Calibri")
        self.accent = rgb(theme.get("accent", "1F4E79"))
        self.ink = rgb(theme.get("ink", "1A1A1A"))
        self.paper = rgb(theme.get("paper", "FFFFFF"))
        self.muted = rgb(theme.get("muted", "666666"))

    def slide(self, number):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = self.paper
        self.current = number
        return s

    def text(self, slide, x, y, w, h, value, size, bold=False, color=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(x, y, w, h)
        frame = box.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = anchor
        frame.margin_left = frame.margin_right = 0
        lines = value if isinstance(value, list) else [value]
        for i, line in enumerate(lines):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.alignment = align
            p.space_after = Pt(size * 0.5)
            run = p.add_run()
            run.text = str(line)
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.name = self.font
            run.font.color.rgb = color or self.ink
        return box

    def headline(self, slide, value):
        if words(value) > MAX_HEADLINE:
            warnings.append(f"slide {self.current}: headline has {words(value)} words (max {MAX_HEADLINE})")
        self.text(slide, MARGIN, Inches(0.6), W - 2 * MARGIN, Inches(1.3), value, 30, bold=True, anchor=MSO_ANCHOR.TOP)
        bar = slide.shapes.add_shape(1, MARGIN, Inches(1.95), Inches(0.9), Emu(38100))
        bar.fill.solid()
        bar.fill.fore_color.rgb = self.accent
        bar.line.fill.background()

    def bullets(self, slide, x, y, w, items, size=22):
        if len(items) > MAX_BULLETS:
            warnings.append(f"slide {self.current}: {len(items)} bullets (max {MAX_BULLETS})")
        for item in items:
            if words(item) > MAX_WORDS:
                warnings.append(f"slide {self.current}: a bullet has {words(item)} words (max {MAX_WORDS}): {str(item)[:40]}")
        self.text(slide, x, y, w, Inches(4.6), [f"•  {i}" for i in items], size)

    def notes(self, slide, value):
        if value:
            slide.notes_slide.notes_text_frame.text = str(value)


def layout_title(d, s, spec):
    d.text(s, MARGIN, Inches(2.4), W - 2 * MARGIN, Inches(1.6), spec.get("title", ""), 44, bold=True, anchor=MSO_ANCHOR.BOTTOM)
    if spec.get("subtitle"):
        d.text(s, MARGIN, Inches(4.2), W - 2 * MARGIN, Inches(1.0), spec["subtitle"], 22, color=d.muted)


def layout_section(d, s, spec):
    d.text(s, MARGIN, Inches(3.0), W - 2 * MARGIN, Inches(1.5), spec.get("title", ""), 38, bold=True, color=d.accent, anchor=MSO_ANCHOR.MIDDLE)


def layout_points(d, s, spec):
    d.headline(s, spec.get("title", ""))
    d.bullets(s, MARGIN, Inches(2.4), W - 2 * MARGIN, spec.get("bullets", []))


def layout_two(d, s, spec):
    d.headline(s, spec.get("title", ""))
    half = (W - 2 * MARGIN - Inches(0.6)) / 2
    for i, side in enumerate(("left", "right")):
        col = spec.get(side, {})
        x = MARGIN + i * (half + Inches(0.6))
        if col.get("heading"):
            d.text(s, x, Inches(2.4), half, Inches(0.6), col["heading"], 24, bold=True, color=d.accent)
        d.bullets(s, x, Inches(3.1), half, col.get("bullets", []), size=20)


def layout_image(d, s, spec):
    d.headline(s, spec.get("title", ""))
    path = spec.get("image", "")
    if not path or not os.path.isfile(path):
        warnings.append(f"slide {d.current}: image not found: {path!r}")
        return
    box_w, box_h = W - 2 * MARGIN, Inches(4.2)
    pic = s.shapes.add_picture(path, MARGIN, Inches(2.3))
    ratio = min(box_w / pic.width, box_h / pic.height)
    pic.width, pic.height = int(pic.width * ratio), int(pic.height * ratio)
    pic.left = int((W - pic.width) / 2)
    if spec.get("caption"):
        d.text(s, MARGIN, Inches(6.65), W - 2 * MARGIN, Inches(0.5), spec["caption"], 14, color=d.muted, align=PP_ALIGN.CENTER)


def layout_table(d, s, spec):
    d.headline(s, spec.get("title", ""))
    header, rows = spec.get("header", []), spec.get("rows", [])
    if len(rows) > 6 or len(header) > 5:
        warnings.append(f"slide {d.current}: table is {len(rows)} rows x {len(header)} columns (max 6 x 5)")
    if not header:
        return
    shape = s.shapes.add_table(len(rows) + 1, len(header), MARGIN, Inches(2.4), W - 2 * MARGIN, Inches(0.5) * (len(rows) + 1))
    for c, label in enumerate(header):
        cell = shape.table.cell(0, c)
        cell.text = str(label)
        cell.fill.solid()
        cell.fill.fore_color.rgb = d.accent
        cell.text_frame.paragraphs[0].runs[0].font.color.rgb = d.paper
        cell.text_frame.paragraphs[0].runs[0].font.bold = True
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row[: len(header)]):
            shape.table.cell(r, c).text = str(value)
    for row in shape.table.rows:
        for cell in row.cells:
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(18)
                    run.font.name = d.font


def layout_quote(d, s, spec):
    d.text(s, Inches(1.4), Inches(2.0), W - Inches(2.8), Inches(3.0), "“" + spec.get("quote", "") + "”", 32, anchor=MSO_ANCHOR.MIDDLE)
    if spec.get("by"):
        d.text(s, Inches(1.4), Inches(5.2), W - Inches(2.8), Inches(0.6), "— " + spec["by"], 20, color=d.muted)


def layout_numbers(d, s, spec):
    d.headline(s, spec.get("title", ""))
    items = spec.get("items", [])[:4]
    if len(spec.get("items", [])) > 4:
        warnings.append(f"slide {d.current}: more than 4 figures; showing 4")
    if not items:
        return
    gap = Inches(0.5)
    each = (W - 2 * MARGIN - gap * (len(items) - 1)) / len(items)
    for i, item in enumerate(items):
        x = MARGIN + i * (each + gap)
        d.text(s, x, Inches(2.8), each, Inches(1.4), item.get("value", ""), 54, bold=True, color=d.accent)
        d.text(s, x, Inches(4.3), each, Inches(1.4), item.get("label", ""), 20, color=d.muted)


LAYOUTS = {"title": layout_title, "section": layout_section, "points": layout_points, "two-column": layout_two,
           "image": layout_image, "table": layout_table, "quote": layout_quote, "numbers": layout_numbers}


def main():
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    spec_path, out_path = sys.argv[1], sys.argv[2]
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    base = os.path.dirname(os.path.abspath(spec_path))
    deck = Deck(spec.get("theme", {}))
    for n, item in enumerate(spec.get("slides", []), start=1):
        kind = item.get("layout", "points")
        if kind not in LAYOUTS:
            warnings.append(f"slide {n}: unknown layout {kind!r}, used points")
            kind = "points"
        if item.get("image") and not os.path.isabs(item["image"]):
            item["image"] = os.path.join(base, item["image"])
        slide = deck.slide(n)
        LAYOUTS[kind](deck, slide, item)
        deck.notes(slide, item.get("notes"))
    deck.prs.save(out_path)
    for line in warnings:
        print("warning:", line)
    print(f"wrote {out_path}: {len(deck.prs.slides)} slides, {len(warnings)} warning(s)")


if __name__ == "__main__":
    main()
